"""Research report discovery and loading helpers."""
from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber

from config import SAMPLE_DIR
from src.knowledge.ov_adapter import store_resource

logger = logging.getLogger(__name__)

RESEARCH_ROOT = SAMPLE_DIR / "附件5：研报数据"
STOCK_METADATA_XLSX = RESEARCH_ROOT / "个股_研报信息.xlsx"
INDUSTRY_METADATA_XLSX = RESEARCH_ROOT / "行业_研报信息.xlsx"


@dataclass(frozen=True)
class ResearchDocument:
    title: str
    paper_path: str
    category: str
    metadata: dict[str, Any]


@dataclass(frozen=True)
class LoadedResearchDocument:
    document: ResearchDocument
    uri: str
    fallback_used: bool = False


def _load_metadata(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return {}
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    mapping: dict[str, dict[str, Any]] = {}
    for row in rows[1:]:
        if not row:
            continue
        record = {headers[idx]: row[idx] for idx in range(min(len(headers), len(row)))}
        title = str(record.get("title") or "").strip()
        if title:
            mapping[title] = record
    return mapping


def discover_research_documents(root_dir: Path = RESEARCH_ROOT) -> list[ResearchDocument]:
    stock_meta = _load_metadata(root_dir / "个股_研报信息.xlsx")
    industry_meta = _load_metadata(root_dir / "行业_研报信息.xlsx")
    documents: list[ResearchDocument] = []
    for category, sub_dir, meta in [("stock", "个股研报", stock_meta), ("industry", "行业研报", industry_meta)]:
        for pdf_path in sorted((root_dir / sub_dir).glob("*.pdf")):
            title = pdf_path.stem
            documents.append(ResearchDocument(title=title, paper_path=str(pdf_path), category=category, metadata=dict(meta.get(title, {}))))
    return documents


def _extract_pdf_text(pdf_path: Path) -> str:
    parts: list[str] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
    return "\n".join(parts)


def _chunk_text(text: str, chunk_size: int = 800) -> list[str]:
    normalized = " ".join(text.split())
    if not normalized:
        return []
    return [normalized[idx:idx + chunk_size] for idx in range(0, len(normalized), chunk_size)]


def _store_fallback_text(client: Any, document: ResearchDocument) -> str:
    pdf_path = Path(document.paper_path)
    chunks = _chunk_text(_extract_pdf_text(pdf_path))
    if not chunks:
        raise RuntimeError(f"No extractable text found in {pdf_path}")
    cache_dir = pdf_path.parent / ".ov_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    txt_path = cache_dir / f"{pdf_path.stem}.txt"
    txt_path.write_text("\n\n".join(chunks), encoding="utf-8")
    return store_resource(client, txt_path)


def _is_already_loaded(client: Any) -> bool:
    """Check if OV already has real research resources (not just system metadata)."""
    try:
        existing = client.ls("/resources")
        return any(item.get("isDir", False) for item in existing)
    except Exception:
        return False


def load_research_documents(client: Any, root_dir: Path = RESEARCH_ROOT) -> list[LoadedResearchDocument]:
    if _is_already_loaded(client):
        logger.info("OpenViking already has indexed resources, skipping re-load.")
        documents = discover_research_documents(root_dir)
        return [LoadedResearchDocument(document=doc, uri=doc.paper_path) for doc in documents]

    loaded: list[LoadedResearchDocument] = []
    for document in discover_research_documents(root_dir):
        try:
            uri = store_resource(client, document.paper_path)
            loaded.append(LoadedResearchDocument(document=document, uri=uri, fallback_used=False))
        except Exception:
            try:
                uri = _store_fallback_text(client, document)
                loaded.append(LoadedResearchDocument(document=document, uri=uri, fallback_used=True))
            except Exception as exc:
                logger.warning("Failed to load research document %s after fallback: %s", document.paper_path, exc)
                continue
    return loaded
